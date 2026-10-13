# -*- coding: utf-8 -*-
"""
SKY Mengde Sammenligning
Sammenligner mengder fra Leveransemodell, Anbud og Omega (K10).

Leveransemodell armering:
  - FOB-kolonne  : direkte sum av FOB_Info/FOB_Mengde per mengdelistepost
  - PE-kolonne   : arm_vekt per PE (PGF_RIB_Info/Vekt fra _RIB_ARM_-rader),
                   tilordnet armering-post via PE-matching mot _RIB_XXX_-rader
                   (identifisert der FOB_Mengdeenhet = 'kg' og posten er
                   ein kjend kg-post i Omega)
"""

import os
import re
import numpy as np
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

INPUT_LEVERANSEMODELL = r"C:\Users\ifi\Downloads\SKY_leveransemodell.csv"
INPUT_ANBUD           = r"C:\Users\ifi\Downloads\SKY_addendum_anbud.csv"
INPUT_OMEGA           = r"C:\Users\ifi\Downloads\SKY_K10_Omega_22052026.xlsx"
OUTPUT_DIR            = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "excel")   # ifi/excel — ikkje i git
OUTPUT_BASE           = "SKY_mengde_sammenligning"

ARM_TOKEN      = "_RIB_ARM_"
RIB_TOKEN      = "_RIB_XXX_"
LP_TOKEN       = "K10-LP"
HOVED_STATUS   = "S5"
BETONG_FAGKAP  = {"05 Betongarbeider", "06 Betongkonstruksjoner, prefab"}

def _hent_modelltype(path):
    """
    Ekstraherer disiplin-koden mellom _SKY_ og _K10 i IFC-filnamnet.
    Eks: '...U_F_RIB_SKY_RIB_ARM_K10-LP020.ifc/...' -> 'RIB_ARM'
    """
    if not path or str(path) == "nan":
        return ""
    for segment in str(path).split("/"):
        if segment.endswith(".ifc"):
            m = re.search(r"_SKY_([A-Z0-9_]+?)_K10", segment)
            if m:
                return m.group(1)
            # Fallback: alt etter _SKY_ til .ifc
            m2 = re.search(r"_SKY_([A-Z0-9_]+?)\.ifc", segment)
            if m2:
                return m2.group(1)
            break
    return ""


# =============================================================================
# Hjelpe-parsere
# =============================================================================

def _parse_verdi(s):
    """Parser norsk desimalkomma-streng til float. Striper asterisker."""
    s = str(s).strip().rstrip("*").strip()
    if not s or s in ("nan", "--", "FOB_Nullpunkt", "= Volume", "False", "True"):
        return np.nan
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return np.nan


def _les_csv(path):
    return pd.read_csv(path, sep=None, engine="python", dtype=str)


def _dropp_ugyldige_kjelder(df, kilde):
    """
    Fjernar rader som ikkje skal vere med i mengdesamanstillinga:
      - DWG-teikningar (CAD-rader med rå NS3420-kodar utan prosjektprefiks)
      - FOB_Nullpunkt-referanseobjekt (origo-anker i BIM-modellen, ikkje mengder)
    """
    if "Path" not in df.columns:
        return df
    n_before = len(df)
    mask = (
        ~df["Path"].str.contains(r"\.dwg", case=False, na=False, regex=True)
        & ~df["Path"].str.contains("FOB_Nullpunkt", na=False)
    )
    df = df[mask].copy()
    n_drop = n_before - len(df)
    if n_drop:
        print(f"   [{kilde}] Droppa {n_drop} rader (DWG-teikningar / FOB_Nullpunkt)")
    return df


def _norm_enhet(e):
    """Normaliser einheit for samanlikning: m³/m² -> m3/m2, lowercase, strip."""
    if e is None or (isinstance(e, float) and np.isnan(e)):
        return ""
    return str(e).strip().lower().replace("³", "3").replace("²", "2")


def _ny_filsti():
    """Generer alltid ny, timestampa filsti so eksisterande filer ikkje vert overskrivne."""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return os.path.join(OUTPUT_DIR, f"{OUTPUT_BASE}_{ts}.xlsx")


# =============================================================================
# Fuzzy-matching av mengdelistepost-koder mot Omega
# =============================================================================

def _norm(code):
    """Kanonisk form: erstatt bindestreker/mellomrom med punktum, lowercase."""
    return re.sub(r"[-\s]+", ".", code.strip()).lower()


def bygg_fuzzy_tabell(omega_posts):
    """Dict: normalisert_kode -> original_omega_kode."""
    return {_norm(p): p for p in omega_posts}


def fuzzy_korriger(code, omega_set, fuzzy_tab):
    """
    Prøver å finne rett Omega-kode for ein (potensielt feil) post-kode.

    Forsøksrekkefølge:
      0. Eksakt match
      1. Mellomrom-splitting: "15.411 1307-11" -> prøv kvart token
      2. Separator-normalisering (- -> .)
      3. Legg til "1" framfor 3-sifra prefiks  (301.xxx -> 1301.xxx)
      4. Kombiner 2+3
      5. Trimm eitt avsluttande null i eit segment (040 -> 04, 100 -> 10)
         + kombiner med prefiks +1
      6. Fjern "1" framfor 4-sifra prefiks (1401.xxx -> 401.xxx)

    Returnerer (korrigert_kode, match_type) der match_type er None ved ingen treff.
    """
    if code in omega_set:
        return code, None

    # 1. Mellomrom-splitting: koden kan vere brote opp i to delar (t.d. "77.19 1307-11")
    #    der den eine delen er prosjektprefiks ("1307-11") og den andre er suffiks ("77.19").
    #    Forsøk å setje dei saman før vi prøver individuelle tokens.
    if " " in code:
        tokens = [t.strip() for t in code.split() if t.strip()]
        if len(tokens) == 2:
            # Identifiser prefiks-token (3-4 siffer etterfulgt av - eller .)
            prefiks_idx = next(
                (i for i, t in enumerate(tokens) if re.match(r"^\d{3,4}[-.]", t)),
                None,
            )
            if prefiks_idx is not None:
                prefiks = tokens[prefiks_idx]
                suffiks = tokens[1 - prefiks_idx]
                for sep in (".", "-"):
                    kombo = f"{prefiks}{sep}{suffiks}"
                    if kombo in omega_set:
                        return kombo, "kombinert mellomrom"
                    kn = _norm(kombo)
                    if kn in fuzzy_tab:
                        return fuzzy_tab[kn], "kombinert mellomrom+separator"
        # Fallback: direkte token-treff
        for token in tokens:
            if token in omega_set:
                return token, "mellomrom"
            tn = _norm(token)
            if tn in fuzzy_tab:
                return fuzzy_tab[tn], "mellomrom+separator"
        # Til slutt: returner tokenet som ser mest ut som ein prosesskode,
        # slik at korriger_foreldrekode kan prøve foreldrekode-matching på det
        for token in reversed(tokens):
            if re.match(r"^\d{3,4}[-.]", token):
                return token, "mellomrom"

    n = _norm(code)

    # 2. Separator-normalisering
    if n in fuzzy_tab:
        return fuzzy_tab[n], "separator"

    # 3+4. Manglar leiande siffer (3-sifra prefiks -> 4-sifra)
    m = re.match(r"^(\d{3})(\..*)", n)
    if m:
        candidate = "1" + m.group(1) + m.group(2)
        if candidate in fuzzy_tab:
            return fuzzy_tab[candidate], "prefiks +1"
        cand_orig = "1" + code
        if cand_orig in omega_set:
            return cand_orig, "prefiks +1"

    # 5. Trimm eitt avsluttande null i eit segment (f.eks. "040"→"04", "100"→"10")
    parts = n.split(".")
    for i in range(1, len(parts)):
        if len(parts[i]) >= 2 and parts[i].endswith("0"):
            trimma = parts.copy()
            trimma[i] = parts[i][:-1]
            candidate = ".".join(trimma)
            if candidate in fuzzy_tab:
                return fuzzy_tab[candidate], "trimma null"
            # Kombiner med prefiks +1
            m3 = re.match(r"^(\d{3})(\..*)", candidate)
            if m3:
                with_prefix = "1" + candidate
                if with_prefix in fuzzy_tab:
                    return fuzzy_tab[with_prefix], "prefiks+trimma null"

    # 6. For mykje leiande siffer (1401.xxx -> 401.xxx)
    m2 = re.match(r"^1(\d{3})(\..*)", n)
    if m2:
        candidate = m2.group(1) + m2.group(2)
        if candidate in fuzzy_tab:
            return fuzzy_tab[candidate], "prefiks -1"

    return code, None


def korriger_poster(df_posts, omega_set, fuzzy_tab, kilde):
    """
    Korrigerer post-koder via fuzzy_korriger.
    Returnerer (korrigert_df, set_av_korrigerte_kodar).
    """
    remap = {}
    for code in df_posts["post"].unique():
        _, mtype = fuzzy_korriger(code, omega_set, fuzzy_tab)
        if mtype:
            korrigert, _ = fuzzy_korriger(code, omega_set, fuzzy_tab)
            remap[code] = (korrigert, mtype)

    if remap:
        print(f"   [{kilde}] Fuzzy-korreksjonar ({len(remap)}):")
        for orig, (ny, mtype) in sorted(remap.items()):
            print(f"     {orig!r:40s} -> {ny!r}  ({mtype})")
        df_posts = df_posts.copy()
        df_posts["post"] = df_posts["post"].map(lambda c: remap[c][0] if c in remap else c)
    else:
        print(f"   [{kilde}] Ingen fuzzy-korreksjonar nodvendig")

    korrigerte = {ny for ny, _ in remap.values()}
    korreksjonar = [
        {"Opphavleg kode": orig, "Korrigert til": ny, "Type korreksjon": mtype}
        for orig, (ny, mtype) in remap.items()
    ]
    return df_posts, korrigerte, korreksjonar


def korriger_foreldrekode(df_posts, omega_set, omega_enhet_map, kilde):
    """
    Foreldrekode-matching: postar som ikkje er i Omega, men som har Omega-barn
    (postar som startar med koden + "."), vert knytt til rett barn via einheit.

    Eks: "1307.01.22.255.1" + enhet "m3" -> "1307.01.22.255.1.1" (om .1.1 har m3 i Omega)

    Krev at df_posts har kolonne 'enhet'.
    """
    if "enhet" not in df_posts.columns:
        return df_posts, set(), []

    # Bygg normalisert oppslag: norm(omega_post) -> omega_post
    norm_to_omega = {_norm(p): p for p in omega_set}

    # Finn alle postar som ikkje er i Omega etter tidlegare fuzzy-korreksjon
    ukjende = [c for c in df_posts["post"].unique() if c not in omega_set]

    # For kvar ukjend kode: finn Omega-barn
    parent_map = {}  # code -> {enhet_str -> omega_child}
    for code in ukjende:
        n = _norm(code) + "."
        barn = [orig for norm_k, orig in norm_to_omega.items() if norm_k.startswith(n)]
        if not barn:
            continue
        enhet_til_barn = {}
        for barn_post in barn:
            e = str(omega_enhet_map.get(barn_post, "")).lower().strip()
            if e not in enhet_til_barn:
                enhet_til_barn[e] = barn_post
        parent_map[code] = enhet_til_barn

    if not parent_map:
        return df_posts, set(), []

    print(f"   [{kilde}] Foreldrekode-matching ({len(parent_map)} postar):")
    for code, barn_dict in sorted(parent_map.items()):
        print(f"     {code!r:40s} -> {list(barn_dict.values())}")

    df_posts = df_posts.copy()
    log = {}  # (orig, korrigert) -> enhet (for dedup)
    drop_idx = []  # indeksar som skal droppast (barn finst allereie i same kjelderad)

    # Originale post-sett per kjelderad - brukt for å sjå om barn allereie er der
    posts_per_row = (
        df_posts.groupby("_row_id")["post"].apply(set).to_dict()
        if "_row_id" in df_posts.columns
        else {}
    )

    def resolve(row):
        code = row["post"]
        if code not in parent_map:
            return code
        barn_dict = parent_map[code]
        enhet = str(row.get("enhet", "")).lower().strip()
        if enhet in barn_dict:
            child = barn_dict[enhet]
        else:
            child = next(iter(barn_dict.values()))
            enhet = ""
        # Om barn allereie er i same kjelderad: drop foreldrekoden, ikkje logg som korreksjon
        row_id = row.get("_row_id")
        if row_id is not None and child in posts_per_row.get(row_id, set()):
            drop_idx.append(row.name)
            return code
        log[(code, child)] = enhet
        return child

    df_posts["post"] = df_posts.apply(resolve, axis=1)
    if drop_idx:
        df_posts = df_posts.drop(index=drop_idx)
        print(f"   [{kilde}] Droppa {len(drop_idx)} foreldrekode-rader (barn finst allereie i same kjelderad)")

    korrigerte = {child for _, child in log}
    korreksjonar = [
        {"Opphavleg kode": orig, "Korrigert til": child,
         "Type korreksjon": f"foreldrekode (einheit: {enhet})" if enhet else "foreldrekode (fallback)"}
        for (orig, child), enhet in log.items()
    ]
    return df_posts, korrigerte, korreksjonar


# =============================================================================
# Ekspansjon av semikolon-separerte post/mengde/einheit
# =============================================================================

def _ekspander_semikol(df, post_kol, mengde_kol, enhet_kol=None, ekstra_kolar=None):
    """
    Ekspanderer semikolon-separerte post/mengde-par til individuelle rader.
    ekstra_kolar: liste med kolonnenamn som kopierast direkte (ikkje semikol-splitta).
    """
    ekstra_kolar = ekstra_kolar or []
    records = []
    for row_id, (_, row) in enumerate(df.iterrows()):
        post_raw = row[post_kol]
        if pd.isna(post_raw):
            continue
        mengde_raw = row[mengde_kol]
        enhet_raw  = row.get(enhet_kol) if enhet_kol else None

        posts   = [p.strip() for p in str(post_raw).split(";")]
        mengdes = [m.strip() for m in str(mengde_raw).split(";")] if pd.notna(mengde_raw) else []
        enheter = [e.strip() for e in str(enhet_raw).split(";")] if (enhet_kol and pd.notna(enhet_raw)) else []

        for i, post in enumerate(posts):
            if not post or post in ("--", ""):
                continue
            m_str = mengdes[i] if i < len(mengdes) else ""
            e_str = enheter[i] if i < len(enheter) else ""
            rec = {"post": post, "mengde": _parse_verdi(m_str), "_row_id": row_id}
            if enhet_kol:
                rec["enhet"]     = _norm_enhet(e_str)
                rec["enhet_raw"] = e_str.strip()
            for k in ekstra_kolar:
                rec[k] = row.get(k, "")
            records.append(rec)

    if not records:
        cols = ["post", "mengde", "_row_id"] + (["enhet", "enhet_raw"] if enhet_kol else []) + ekstra_kolar
        return pd.DataFrame(columns=cols)
    return pd.DataFrame(records)


def _dedup_innan_kjelderad(exp, kilde):
    """
    Droppar duplikat (_row_id, post) etter korreksjonar.
    Stoppar dobbeltteljing når ein foreldrekode korrigerast til ein barn-kode
    som allereie finst i same kjelderad (semikolon-listen).
    """
    if "_row_id" not in exp.columns:
        return exp
    n_before = len(exp)
    exp = exp.drop_duplicates(subset=["_row_id", "post"], keep="first")
    n_drop = n_before - len(exp)
    if n_drop:
        print(f"   [{kilde}] Droppa {n_drop} duplikatrader (same post i same kjelderad etter korreksjon)")
    return exp


# =============================================================================
# Leveransemodell – FOB_Mengde-kolonne (standard)
# =============================================================================

def _filtrer_einheit_mismatch(exp, omega_enhet_map, kilde):
    """
    Fjernar rader der FOB-einheit ikkje stemmer med Omega-einheit (begge normaliserte).
    Døme: betong-postar (m³) skal ikkje få kg-rader summert inn.
    """
    if not omega_enhet_map or "enhet" not in exp.columns:
        return exp
    n_before = len(exp)
    mask = exp.apply(
        lambda row: (
            True
            if omega_enhet_map.get(row["post"]) is None or not str(row.get("enhet", "")).strip()
            else str(row.get("enhet", "")).strip() == omega_enhet_map[row["post"]]
        ),
        axis=1,
    )
    exp = exp[mask].copy()
    n_filt = n_before - len(exp)
    if n_filt:
        print(f"   [{kilde}] Filtrert {n_filt} rader med einheit-mismatch mot Omega")
    return exp


def _samle_fob_enhet(exp):
    """Returnerer dict post -> vanlegaste rå-einheit (bevarer original casing for display)."""
    col = "enhet_raw" if "enhet_raw" in exp.columns else "enhet"
    if col not in exp.columns:
        return {}
    result = {}
    for post, grp in exp.groupby("post"):
        counts = grp[col].dropna().value_counts()
        if len(counts) and counts.index[0]:
            result[post] = counts.index[0]
    return result


def load_leveransemodell_fob(path, omega_set, fuzzy_tab, omega_enhet_map):
    df = _les_csv(path)
    df = _dropp_ugyldige_kjelder(df, "Leveransemodell")
    df_valid = df[df["FOB_Info/FOB_Mengdelistepost"].notna()].copy()
    df_valid["_type"] = df_valid["Path"].apply(_hent_modelltype)
    exp = _ekspander_semikol(
        df_valid,
        "FOB_Info/FOB_Mengdelistepost",
        "FOB_Info/FOB_Mengde",
        enhet_kol="FOB_Info/FOB_Mengdeenhet",
        ekstra_kolar=["_type"],
    )
    exp = exp.dropna(subset=["mengde"])
    exp, k1, log1 = korriger_poster(exp, omega_set, fuzzy_tab, "Leveransemodell")
    exp, k2, log2 = korriger_foreldrekode(exp, omega_set, omega_enhet_map, "Leveransemodell")
    exp = _dedup_innan_kjelderad(exp, "Leveransemodell")
    exp = _filtrer_einheit_mismatch(exp, omega_enhet_map, "Leveransemodell")
    fob_enhet_map = _samle_fob_enhet(exp)
    summary = (
        exp.groupby("post")
        .agg(
            mengde=("mengde", "sum"),
            Kilde=("_type", lambda x: " / ".join(sorted({v for v in x if v}))),
        )
        .reset_index()
        .rename(columns={"post": "Mengdelistepost", "mengde": "Mengde Leveransemodell"})
    )
    korreksjonar = [{**r, "Kjelde": "Leveransemodell"} for r in log1 + log2]
    return summary, k1 | k2, korreksjonar, fob_enhet_map


# =============================================================================
# Leveransemodell – PE-kolonne (armering via Produksjonsenhet-matching)
# =============================================================================

def load_leveransemodell_pe(path, omega_kg_posts):
    """
    Beregner armeringsvekt per mengdelistepost via PE-matching:
      1. ARM-rader (_RIB_ARM_): sum PGF_RIB_Info/Vekt per PE (LM-korrigert)
      2. RIB-rader (_RIB_XXX_): finn armering-post (der FOB_Mengdeenhet='kg'
         OG posten finst i omega_kg_posts)
      3. Koble: post -> PE -> arm_vekt

    omega_kg_posts: set av Omega-postar med einheit 'kg'.
    """
    df = _les_csv(path)
    df = _dropp_ugyldige_kjelder(df, "Leveransemodell PE")

    # --- ARM-vekt per PE ---
    arm_raw = df[
        df["Path"].str.contains(LP_TOKEN,  na=False)
        & df["Path"].str.contains(ARM_TOKEN, na=False)
        & df["FOB_Info/FOB_Produksjonsenhet"].notna()
        & df["PGF_RIB_Info/Vekt"].notna()
    ].copy()
    arm_raw["vekt"] = pd.to_numeric(arm_raw["PGF_RIB_Info/Vekt"], errors="coerce")

    # LM-korreksjon: Formkode med "LM" tel berre 7.2% av vekta
    formkode_col = "PGF_RIB_Info/Formkode"
    if formkode_col in arm_raw.columns:
        er_lm = arm_raw[formkode_col].astype(str).str.contains("LM", case=False, na=False)
        lm_count = er_lm.sum()
        lm_vekt  = arm_raw.loc[er_lm, "vekt"].sum()
        print(f"   LM-deteksjon: {lm_count} jern, total LM-vekt = {lm_vekt:.2f} kg")
        arm_raw["vekt_eff"] = np.where(er_lm, arm_raw["vekt"] * 0.072, arm_raw["vekt"])
    else:
        print("   [ADVARSEL] PGF_RIB_Info/Formkode ikkje funnet - LM-korreksjon hoppast over")
        arm_raw["vekt_eff"] = arm_raw["vekt"]

    arm_raw["FOB_Info/FOB_Produksjonsenhet"] = (
        arm_raw["FOB_Info/FOB_Produksjonsenhet"].astype(str).str.split(",")
    )
    arm_exp = arm_raw.explode("FOB_Info/FOB_Produksjonsenhet")
    arm_exp["FOB_Info/FOB_Produksjonsenhet"] = arm_exp["FOB_Info/FOB_Produksjonsenhet"].str.strip()
    arm_exp = arm_exp[arm_exp["FOB_Info/FOB_Produksjonsenhet"] != ""]
    arm_per_pe = arm_exp.groupby("FOB_Info/FOB_Produksjonsenhet")["vekt_eff"].sum().to_dict()
    print(f"   PE-analyse: {len(arm_per_pe)} unike PE-ar med ARM-vekt (inkl. LM-korreksjon)")

    # --- Finn (PE -> armering-post) frå RIB-betong-rader ---
    rib = df[
        df["Path"].str.contains(LP_TOKEN,  na=False)
        & df["Path"].str.contains(RIB_TOKEN, na=False)
        & df["FOB_Info/FOB_Mengdelistepost"].notna()
        & df["FOB_Info/FOB_Produksjonsenhet"].notna()
        & df["PGF_Info/PGF_Fagkapittel"].isin(BETONG_FAGKAP)
        & ~df["PGF_Info/PGF_Mengdetype"].str.contains("fuge", case=False, na=False)
    ].copy()

    pe_to_post = {}
    for _, row in rib.iterrows():
        pe_raw   = str(row["FOB_Info/FOB_Produksjonsenhet"])
        pe_list  = [p.strip() for p in pe_raw.split(",") if p.strip()]
        posts    = [p.strip() for p in str(row["FOB_Info/FOB_Mengdelistepost"]).split(";")]
        enheter  = [e.strip().lower() for e in str(row.get("FOB_Info/FOB_Mengdeenhet", "")).split(";")]
        arm_post = None
        for post, enhet in zip(posts, enheter):
            if enhet == "kg" and post:
                if not omega_kg_posts or post in omega_kg_posts:
                    arm_post = post
                    break
        if arm_post is None:
            continue
        for pe in pe_list:
            if pe and pe not in pe_to_post:
                pe_to_post[pe] = arm_post

    print(f"   PE-analyse: {len(pe_to_post)} PE-ar med identifisert armering-post")

    records = [{"post": post, "vekt": arm_per_pe[pe]}
               for pe, post in pe_to_post.items() if pe in arm_per_pe]

    if not records:
        return pd.DataFrame(columns=["Mengdelistepost", "Mengde Lev (PE)"])

    return (
        pd.DataFrame(records)
        .groupby("post")["vekt"]
        .sum()
        .reset_index()
        .rename(columns={"post": "Mengdelistepost", "vekt": "Mengde Lev (PE)"})
    )


# =============================================================================
# Anbud – FOB_Mengde
# =============================================================================

def load_anbud(path, omega_set, fuzzy_tab, omega_enhet_map):
    df = _les_csv(path)
    df = _dropp_ugyldige_kjelder(df, "Anbud")
    df_valid = df[df["FOB_Info/FOB_Mengdelistepost"].notna()].copy()
    df_valid["_type"] = df_valid["Path"].apply(_hent_modelltype)
    exp = _ekspander_semikol(
        df_valid,
        "FOB_Info/FOB_Mengdelistepost",
        "FOB_Info/FOB_Mengde",
        enhet_kol="FOB_Info/FOB_Mengdeenhet",
        ekstra_kolar=["_type"],
    )
    exp = exp.dropna(subset=["mengde"])
    exp, k1, log1 = korriger_poster(exp, omega_set, fuzzy_tab, "Anbud")
    exp, k2, log2 = korriger_foreldrekode(exp, omega_set, omega_enhet_map, "Anbud")
    exp = _dedup_innan_kjelderad(exp, "Anbud")
    exp = _filtrer_einheit_mismatch(exp, omega_enhet_map, "Anbud")
    fob_enhet_map = _samle_fob_enhet(exp)
    summary = (
        exp.groupby("post")
        .agg(
            mengde=("mengde", "sum"),
            Kilde_anb=("_type", lambda x: " / ".join(sorted({v for v in x if v}))),
        )
        .reset_index()
        .rename(columns={"post": "Mengdelistepost", "mengde": "Mengde Anbud"})
    )
    korreksjonar = [{**r, "Kjelde": "Anbud"} for r in log1 + log2]
    return summary, k1 | k2, korreksjonar, fob_enhet_map


# =============================================================================
# Omega
# =============================================================================

def bygg_pe_flytting(anb_path, lev_path, omega_tittel_map=None):
    """
    Identifiserer BIM-element der mengdelisteposten har endra seg mellom
    Anbud og Leveransemodell, og sjekkar om total-mengde per einheit er bevart.

    Brukar IFC-element-ID frå Path (siste numeriske segment, t.d. '7151549')
    fordi Anbud-CSV-en ikkje har FOB_Produksjonsenhet utfylt.
    """
    def _element_id(path):
        """
        Hentar ut IFC-element-ID frå Path. ID-en er 4+ siffer etter siste kolon,
        eventuelt etterfulgt av ein duplikat-suffiks ':N' (1-3 siffer).
        Eks: 'Floor:Dekke 400 B45:1745447:2' -> '1745447'
        """
        if not isinstance(path, str):
            return None
        last = path.rstrip("/").split("/")[-1]
        m = re.search(r":(\d{4,})(?::\d{1,3})?$", last)
        return m.group(1) if m else None

    def _les_element(path, label):
        df = _les_csv(path)
        df = _dropp_ugyldige_kjelder(df, label)
        df = df[df["FOB_Info/FOB_Mengdelistepost"].notna()].copy()
        df["_eid"] = df["Path"].apply(_element_id)
        df = df[df["_eid"].notna()].copy()
        records = []
        for _, row in df.iterrows():
            eid = row["_eid"]
            posts = [p.strip() for p in str(row["FOB_Info/FOB_Mengdelistepost"]).split(";")]
            meng  = str(row.get("FOB_Info/FOB_Mengde", "")).split(";")
            enh   = str(row.get("FOB_Info/FOB_Mengdeenhet", "")).split(";")
            for i, p in enumerate(posts):
                if not p or p in ("--", ""):
                    continue
                m = _parse_verdi(meng[i] if i < len(meng) else "")
                if np.isnan(m):
                    continue
                e = _norm_enhet(enh[i] if i < len(enh) else "")
                records.append({"PE": eid, "post": p, "enhet": e, "mengde": m})
        return pd.DataFrame(records)
    _les_pe = _les_element  # behald gamle namnet på funksjonen som vert kalla

    print("Bygger PE-flytting-analyse ...")
    anb = _les_pe(anb_path, "Anbud-PE")
    lev = _les_pe(lev_path, "Lev-PE")
    print(f"   {anb['PE'].nunique()} PE i Anbud, {lev['PE'].nunique()} PE i Lev")

    anb_agg = anb.groupby(["PE", "post", "enhet"], as_index=False)["mengde"].sum()
    lev_agg = lev.groupby(["PE", "post", "enhet"], as_index=False)["mengde"].sum()

    anb_per_pe = anb_agg.groupby("PE")
    lev_per_pe = lev_agg.groupby("PE")
    common_pes = sorted(set(anb_agg["PE"]) & set(lev_agg["PE"]))

    def _kategori(e):
        return {
            "m3": "Betong", "kg": "Armering", "m2": "Form/Areal",
            "m": "Lengde", "stk": "Antall", "stk.": "Antall",
        }.get((e or "").lower(), e or "—")

    tittel_map = omega_tittel_map or {}
    def _tittel(post):
        return tittel_map.get(post, "")

    def _type_sig(post):
        """
        Tittel-signatur for paring: tekst etter siste ' - ' (alt skifte til upper),
        t.d. 'Bjelke ... - ARMERING MED KAMSTENGER' -> 'ARMERING MED KAMSTENGER'.
        Postar med same signatur skal parast.
        """
        t = _tittel(post)
        if not t:
            return ""
        parts = str(t).split(" - ")
        return parts[-1].strip().upper()

    def _legg_par(rows, kategori, pe, enhet, ap, am, lp, lm):
        """Bygg ein flytting-rad. ap/lp kan vere None (rein add/fjern)."""
        a_sum = am if am is not None else 0
        l_sum = lm if lm is not None else 0
        if ap is None and lp is None:
            return
        if ap == lp and a_sum == l_sum:
            return  # ingen endring
        diff_pct = ((l_sum - a_sum) / a_sum * 100) if a_sum else None
        bevart = a_sum > 0 and diff_pct is not None and abs(diff_pct) <= 5
        rows.append({
            "Kategori":         kategori,
            "Produksjonsenhet": pe,
            "Einheit":          enhet,
            "Anbud post":       ap or "—",
            "Anbud tittel":     _tittel(ap) if ap else "—",
            "Anbud mengde":     am,
            "Lev post":         lp or "—",
            "Lev tittel":       _tittel(lp) if lp else "—",
            "Lev mengde":       lm,
            "Avvik %":          diff_pct,
            "Mengde bevart":    "JA" if bevart else "NEI",
        })

    rows = []
    for pe in common_pes:
        a = anb_per_pe.get_group(pe)
        l = lev_per_pe.get_group(pe)
        if set(zip(a["post"], a["enhet"])) == set(zip(l["post"], l["enhet"])):
            continue
        for enhet in sorted(set(a["enhet"]) | set(l["enhet"])):
            a_e = a[a["enhet"] == enhet]
            l_e = l[l["enhet"] == enhet]
            a_items = list(zip(a_e["post"].tolist(), a_e["mengde"].tolist()))
            l_items = list(zip(l_e["post"].tolist(), l_e["mengde"].tolist()))
            kat = _kategori(enhet)

            # Grupper etter tittel-signatur
            from collections import defaultdict
            a_by_type = defaultdict(list)
            l_by_type = defaultdict(list)
            for p, m in a_items:
                a_by_type[_type_sig(p)].append((p, m))
            for p, m in l_items:
                l_by_type[_type_sig(p)].append((p, m))

            for sig in sorted(set(a_by_type) | set(l_by_type)):
                al = a_by_type.get(sig, [])
                ll = l_by_type.get(sig, [])
                # 1:1 → direkte par
                if len(al) == 1 and len(ll) == 1:
                    (ap, am), (lp, lm) = al[0], ll[0]
                    _legg_par(rows, kat, pe, enhet, ap, am, lp, lm)
                # Berre Anbud → fjerna postar
                elif al and not ll:
                    for ap, am in al:
                        _legg_par(rows, kat, pe, enhet, ap, am, None, None)
                # Berre Lev → nye postar
                elif ll and not al:
                    for lp, lm in ll:
                        _legg_par(rows, kat, pe, enhet, None, None, lp, lm)
                # n:m (same tittel-signatur, fleire postar på begge sider)
                # → par etter sortert rekkefølge, rest blir add/fjern
                else:
                    al_s = sorted(al)
                    ll_s = sorted(ll)
                    for i in range(max(len(al_s), len(ll_s))):
                        ap, am = al_s[i] if i < len(al_s) else (None, None)
                        lp, lm = ll_s[i] if i < len(ll_s) else (None, None)
                        _legg_par(rows, kat, pe, enhet, ap, am, lp, lm)
    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values(["Kategori", "Mengde bevart", "Produksjonsenhet"],
                            ascending=[True, False, True]).reset_index(drop=True)
    print(f"   {len(df)} parings-rader (ein post/flytting per rad)")
    return df


def load_omega(path):
    df_raw = pd.read_excel(path, sheet_name="BOQNT", header=None)
    cols = df_raw.iloc[2].tolist()
    df   = df_raw.iloc[3:].copy()
    df.columns = cols
    df = df.reset_index(drop=True)

    navn_col = "Navn (Prosesskode)"
    gjeld    = "Gjeldende"
    ep_col   = "Enhetspris"

    df = df[df[navn_col].notna()].copy()
    df[gjeld]  = pd.to_numeric(df[gjeld],  errors="coerce")
    df[ep_col] = pd.to_numeric(df[ep_col], errors="coerce")

    return (
        df.groupby(navn_col)
        .agg(
            Kontraktsmengde=(gjeld,    "sum"),
            Enhetspris      =(ep_col,   "first"),
            Tittel          =("Tittel", "first"),
            Enhet           =("Enhet",  "first"),
        )
        .reset_index()
        .rename(columns={navn_col: "Mengdelistepost"})
    )


# =============================================================================
# Excel-skriving
# =============================================================================

HEADERS = [
    "Mengdelistepost",
    "Tittel",
    "Enhet",
    "Mengde Anbud",
    "Mengde Leveransemodell",
    "Kilde",
    "Mengde Lev (PE)",
    "Kontraktsmengde",
    "Enhetspris",
    "Endret",
]
COL_WIDTHS = [32, 40, 8, 20, 24, 22, 18, 20, 14, 8]

BETONG_HEADERS = [
    "Mengdelistepost",
    "Tittel",
    "Enhet",
    "Kilde",
    "Mengde Anbud",
    "Mengde Leveransemodell",
    "Mengde Lev (PE)",
    "Mengde (effektiv)",
    "Kontraktsmengde",
    "Enhetspris",
]
BETONG_WIDTHS = [32, 40, 8, 22, 20, 24, 18, 20, 20, 14]


def write_excel(df, output_path, endrete_poster, korreksjonar, pe_flytting=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mengde Sammenligning"

    HDR_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    HDR_FILL   = PatternFill("solid", start_color="1F3864")
    DATA_FONT  = Font(name="Arial", size=10)
    ALT_FILL   = PatternFill("solid", start_color="DCE6F1")
    ENDR_FILL  = PatternFill("solid", start_color="FFE699")  # gul for korrigerte rader
    CENTER     = Alignment(horizontal="center", vertical="center")
    LEFT       = Alignment(horizontal="left",   vertical="center")
    thin       = Side(border_style="thin", color="B8CCE4")
    border     = Border(left=thin, right=thin, top=thin, bottom=thin)

    num_fmt = "#,##0.00;(#,##0.00);-"
    fmts    = [None, None, None, num_fmt, num_fmt, None, num_fmt, num_fmt, num_fmt, None]

    for c_idx, (hdr, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=c_idx, value=hdr)
        cell.font      = HDR_FONT
        cell.fill      = HDR_FILL
        cell.alignment = CENTER
        cell.border    = border
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    ws.row_dimensions[1].height = 20

    for r_idx, (_, row) in enumerate(df.iterrows(), start=2):
        er_endret = row["Mengdelistepost"] in endrete_poster
        std_fill  = ENDR_FILL if er_endret else (ALT_FILL if r_idx % 2 == 0 else None)

        for c_idx, (key, fmt) in enumerate(zip(HEADERS, fmts), start=1):
            if key == "Endret":
                v = "x" if er_endret else None
            else:
                val = row[key]
                v   = None if (val is None or (isinstance(val, float) and np.isnan(val))) else val
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font      = DATA_FONT
            cell.border    = border
            cell.alignment = LEFT if key in {"Mengdelistepost", "Tittel", "Kilde"} else CENTER
            if fmt:
                cell.number_format = fmt
            if std_fill:
                cell.fill = std_fill

    ws.freeze_panes = "A2"

    last_row = len(df) + 1
    last_col = get_column_letter(len(HEADERS))
    tabell = Table(displayName="MengdeSammenligning", ref=f"A1:{last_col}{last_row}")
    tabell.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(tabell)

    # --- Ark 2: Betong og ARM ---
    ANBUD_FILL = PatternFill("solid", start_color="FCE4D6")  # oransje = mengde frå anbud
    df_ba = df[df["Kilde"].fillna("").str.contains("RIB", na=False)].copy()
    if not df_ba.empty:
        def _effektiv(row):
            lev = row.get("Mengde Leveransemodell")
            pe  = row.get("Mengde Lev (PE)")
            anb = row.get("Mengde Anbud")
            if pd.notna(lev):
                return lev, False
            if pd.notna(pe):
                return pe, False
            if pd.notna(anb):
                return anb, True
            return None, False
        eff = df_ba.apply(_effektiv, axis=1)
        df_ba["Mengde (effektiv)"] = [e[0] for e in eff]
        df_ba["_anbud_fallback"]   = [e[1] for e in eff]

        ws_ba = wb.create_sheet("Betong og ARM")
        ba_fmts = [None, None, None, None, num_fmt, num_fmt, num_fmt, num_fmt, num_fmt, num_fmt]

        for c_idx, (hdr, width) in enumerate(zip(BETONG_HEADERS, BETONG_WIDTHS), 1):
            cell = ws_ba.cell(row=1, column=c_idx, value=hdr)
            cell.font      = HDR_FONT
            cell.fill      = HDR_FILL
            cell.alignment = CENTER
            cell.border    = border
            ws_ba.column_dimensions[get_column_letter(c_idx)].width = width
        ws_ba.row_dimensions[1].height = 20

        for r_idx, (_, row) in enumerate(df_ba.iterrows(), start=2):
            fallback = bool(row["_anbud_fallback"])
            row_fill = ANBUD_FILL if fallback else (ALT_FILL if r_idx % 2 == 0 else None)
            for c_idx, (key, fmt) in enumerate(zip(BETONG_HEADERS, ba_fmts), 1):
                val = row.get(key)
                v = None if (val is None or (isinstance(val, float) and np.isnan(val))) else val
                cell = ws_ba.cell(row=r_idx, column=c_idx, value=v)
                cell.font      = DATA_FONT
                cell.border    = border
                cell.alignment = LEFT if key in {"Mengdelistepost", "Tittel", "Kilde"} else CENTER
                if fmt:
                    cell.number_format = fmt
                if row_fill:
                    cell.fill = row_fill

        ws_ba.freeze_panes = "A2"
        last_ba = get_column_letter(len(BETONG_HEADERS))
        tabell_ba = Table(displayName="BetongOgARM", ref=f"A1:{last_ba}{len(df_ba)+1}")
        tabell_ba.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws_ba.add_table(tabell_ba)

    # --- Ark 3: Korreksjonar ---
    if korreksjonar:
        df_k = pd.DataFrame(korreksjonar)[
            ["Kjelde", "Opphavleg kode", "Korrigert til", "Type korreksjon"]
        ].drop_duplicates().sort_values(["Kjelde", "Type korreksjon", "Opphavleg kode"])

        ws2 = wb.create_sheet("Korreksjonar")
        K_HEADERS = ["Kjelde", "Opphavleg kode", "Korrigert til", "Type korreksjon"]
        K_WIDTHS  = [18, 40, 40, 36]

        for c_idx, (hdr, width) in enumerate(zip(K_HEADERS, K_WIDTHS), 1):
            cell = ws2.cell(row=1, column=c_idx, value=hdr)
            cell.font      = HDR_FONT
            cell.fill      = HDR_FILL
            cell.alignment = CENTER
            cell.border    = border
            ws2.column_dimensions[get_column_letter(c_idx)].width = width
        ws2.row_dimensions[1].height = 20

        for r_idx, (_, row) in enumerate(df_k.iterrows(), start=2):
            row_fill = ALT_FILL if r_idx % 2 == 0 else None
            for c_idx, key in enumerate(K_HEADERS, 1):
                cell = ws2.cell(row=r_idx, column=c_idx, value=row[key])
                cell.font      = DATA_FONT
                cell.border    = border
                cell.alignment = LEFT
                if row_fill:
                    cell.fill = row_fill

        ws2.freeze_panes = "A2"
        last_k = get_column_letter(len(K_HEADERS))
        tabell_k = Table(displayName="Korreksjonar", ref=f"A1:{last_k}{len(df_k)+1}")
        tabell_k.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium7", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws2.add_table(tabell_k)

    # --- Ark 4: PE-flytting ---
    if pe_flytting is not None and not pe_flytting.empty:
        BEVART_FILL   = PatternFill("solid", start_color="C6EFCE")  # grønn = mengde bevart ±5%
        IKKJEBEVART_F = PatternFill("solid", start_color="FFC7CE")  # rosa  = mengde endra
        ws4 = wb.create_sheet("PE-flytting")
        PE_HEADERS = ["Kategori", "Produksjonsenhet", "Einheit",
                      "Anbud post", "Anbud tittel", "Anbud mengde",
                      "Lev post",   "Lev tittel",   "Lev mengde",
                      "Avvik %", "Mengde bevart"]
        PE_WIDTHS  = [14, 22, 8, 36, 44, 14, 36, 44, 14, 10, 14]
        PE_FMTS    = [None, None, None, None, None, num_fmt, None, None, num_fmt, '0.0"%"', None]

        for c_idx, (hdr, width) in enumerate(zip(PE_HEADERS, PE_WIDTHS), 1):
            cell = ws4.cell(row=1, column=c_idx, value=hdr)
            cell.font      = HDR_FONT
            cell.fill      = HDR_FILL
            cell.alignment = CENTER
            cell.border    = border
            ws4.column_dimensions[get_column_letter(c_idx)].width = width
        ws4.row_dimensions[1].height = 20

        for r_idx, (_, row) in enumerate(pe_flytting.iterrows(), start=2):
            row_fill = BEVART_FILL if row["Mengde bevart"] == "JA" else IKKJEBEVART_F
            for c_idx, (key, fmt) in enumerate(zip(PE_HEADERS, PE_FMTS), 1):
                val = row.get(key)
                v = None if (val is None or (isinstance(val, float) and np.isnan(val))) else val
                cell = ws4.cell(row=r_idx, column=c_idx, value=v)
                cell.font      = DATA_FONT
                cell.border    = border
                cell.alignment = LEFT if key in {"Produksjonsenhet", "Anbud post", "Anbud tittel",
                                                  "Lev post", "Lev tittel", "Kategori"} else CENTER
                if fmt:
                    cell.number_format = fmt
                cell.fill = row_fill

        ws4.freeze_panes = "A2"
        last_pe = get_column_letter(len(PE_HEADERS))
        tabell_pe = Table(displayName="PEFlytting", ref=f"A1:{last_pe}{len(pe_flytting)+1}")
        tabell_pe.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium6", showFirstColumn=False,
            showLastColumn=False, showRowStripes=False, showColumnStripes=False,
        )
        ws4.add_table(tabell_pe)

    wb.save(output_path)
    print(f"Lagret: {output_path}")


# =============================================================================
# Hovudflyt
# =============================================================================

def main():
    print("Laster Omega ...")
    omega = load_omega(INPUT_OMEGA)
    print(f"  {len(omega)} unike mengdelisteposter")

    omega_set      = set(omega["Mengdelistepost"].astype(str))
    fuzzy_tab      = bygg_fuzzy_tabell(omega_set)
    omega_enhet_map = {p: _norm_enhet(e) for p, e in zip(omega["Mengdelistepost"], omega["Enhet"])}

    # Kanonisk visningsformat for einheit (m³, m², KG osv. som Omega skriv det)
    omega_enhet_kanonisk = {}
    for e in omega["Enhet"].dropna().unique():
        omega_enhet_kanonisk[_norm_enhet(e)] = str(e).strip()

    omega_kg_posts = {p for p, e in omega_enhet_map.items() if e == "kg"}
    print(f"  {len(omega_kg_posts)} Omega-postar med kg-einheit (armeringspostar)")

    print("Laster Leveransemodell (FOB_Mengde) ...")
    lev_fob, k_lev, korr_lev, fob_enhet_lev = load_leveransemodell_fob(INPUT_LEVERANSEMODELL, omega_set, fuzzy_tab, omega_enhet_map)
    print(f"  {len(lev_fob)} unike mengdelisteposter")

    print("Laster Leveransemodell (PE-matching) ...")
    lev_pe = load_leveransemodell_pe(INPUT_LEVERANSEMODELL, omega_kg_posts)
    print(f"  {len(lev_pe)} armering-poster via PE")

    print("Laster Anbud ...")
    anb, k_anb, korr_anb, fob_enhet_anb = load_anbud(INPUT_ANBUD, omega_set, fuzzy_tab, omega_enhet_map)
    print(f"  {len(anb)} unike mengdelisteposter")

    # Kombiner FOB-einheit-oppslag: leveransemodell har prioritet
    fob_enhet_fallback = {**fob_enhet_anb, **fob_enhet_lev}

    endrete_poster = k_lev | k_anb
    alle_korreksjonar = korr_lev + korr_anb
    print(f"  {len(endrete_poster)} unike postkoder vart korrigert for Omega-matching")

    print("Slar saman ...")
    merged = (
        anb
        .merge(lev_fob, on="Mengdelistepost", how="outer")
        .merge(lev_pe,  on="Mengdelistepost", how="outer")
        .merge(omega,   on="Mengdelistepost", how="outer")
    )
    merged = merged.sort_values(
        "Mengdelistepost", key=lambda s: s.str.lower()
    ).reset_index(drop=True)

    # Fyll Kilde frå anbud-modelltype når Leveransemodellen ikkje har data
    if "Kilde_anb" in merged.columns:
        def _fill_kilde(row):
            k = row.get("Kilde")
            if pd.isna(k) or str(k).strip() in ("", "nan"):
                return row.get("Kilde_anb")
            return k
        merged["Kilde"] = merged.apply(_fill_kilde, axis=1)
        merged = merged.drop(columns=["Kilde_anb"])

    # Fyll manglande Enhet frå FOB_Mengdeenhet, og normaliser til Omega-format
    def _fill_enhet(row):
        e = row.get("Enhet")
        if pd.isna(e) or str(e).strip() in ("", "nan"):
            e = fob_enhet_fallback.get(row["Mengdelistepost"])
        if e is None or (isinstance(e, float) and np.isnan(e)) or str(e).strip() in ("", "nan"):
            return None
        return omega_enhet_kanonisk.get(_norm_enhet(e), str(e).strip())
    merged["Enhet"] = merged.apply(_fill_enhet, axis=1)

    # Legg til RIB_ARM i Kilde for postar med PE-utrekna armeringsvekt
    has_pe = merged["Mengde Lev (PE)"].notna()
    def _kilde_med_arm(kilde_val):
        k = "" if pd.isna(kilde_val) or str(kilde_val).strip() in ("", "nan") else str(kilde_val).strip()
        if "RIB_ARM" in k:
            return k
        return (k + " / RIB_ARM") if k else "RIB_ARM"
    merged.loc[has_pe, "Kilde"] = merged.loc[has_pe, "Kilde"].apply(_kilde_med_arm)

    # Behold output-kolonnar (Kilde er no del av HEADERS)
    data_headers = [h for h in HEADERS if h != "Endret"]
    for h in data_headers:
        if h not in merged.columns:
            merged[h] = None
    merged = merged[data_headers]
    print(f"  Totalt {len(merged)} unike mengdelisteposter")

    omega_tittel_map = dict(zip(omega["Mengdelistepost"].astype(str), omega["Tittel"].astype(str)))
    pe_flytting = bygg_pe_flytting(INPUT_ANBUD, INPUT_LEVERANSEMODELL, omega_tittel_map)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output_path = _ny_filsti()
    print("Skriv Excel ...")
    write_excel(merged, output_path, endrete_poster, alle_korreksjonar, pe_flytting)


if __name__ == "__main__":
    main()
