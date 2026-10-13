"""
Generer korrigert Omfar_Analyse fra VEK og LYA mengdeuttak-resultater.
Les LM_Omfar fra begge resultfiler og beregn korrekte omfar-statistikker.
"""

import math
import os
import warnings
from datetime import datetime

import pandas as pd

warnings.filterwarnings("ignore")

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# === KILDEFILER ===
# excel-mappa i ifi (eitt nivå over tools/) — ikkje i git
_IFI = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "excel")
os.makedirs(_IFI, exist_ok=True)
VEK_FILE = os.path.join(_IFI, "VEK_resultat_mengdeuttak_Omega_total_K7_2302_20260313_125218.xlsx")
LYA_FILE = os.path.join(_IFI, "LYA_resultat_mengdeuttak_Omega_total_K7_2302_20260313_130045.xlsx")

ts = datetime.now().strftime("%Y%m%d_%H%M%S")
OUTPUT_FILE = os.path.join(_IFI, f"Omfar_Analyse_{ts}.xlsx")

# === LES DATA ===
print("[INFO] Leser VEK LM_Omfar ...")
df_vek = pd.read_excel(VEK_FILE, sheet_name="LM_Omfar")

# Fjern eventuelle hjelpekolonner (f.eks. Kolonne1)
df_vek = df_vek[[c for c in df_vek.columns if not c.startswith("Kolonne")]].copy()

print("[INFO] Leser LYA LM_Omfar ...")
df_lya = pd.read_excel(LYA_FILE, sheet_name="LM_Omfar")
df_lya = df_lya[[c for c in df_lya.columns if not c.startswith("Kolonne")]].copy()

# Kolonnenavn-snarveier
C_VEKT      = "Vekt (kg)"
C_DIAM      = "Diameter (mm)"
C_LENGDE    = "Lengde (m)"
C_DEKKER    = "Dekker >250mm"
C_ANT_GEN   = "Antall omfar [=ceil((L-12)/(12-omfar_L))]"
C_VEKT_GEN  = "Vekt inkl. omfar (kg)"
C_MENGDE_GEN = "Mengde omfar (m) [=n*omfar_L]"
C_ANDEL_GEN  = "Omfar andel (%) [=(vekt_inkl-vekt)/vekt*100]"

# === BEREGN STATISTIKK PER KILDE ===

def stats_for(df, label):
    ant_total = len(df)
    mask_omfar = df[C_ANT_GEN].fillna(0) > 0
    ant_med_omfar = mask_omfar.sum()
    vekt_uten = df[C_VEKT].sum()
    vekt_inkl_gen = df[C_VEKT_GEN].sum()
    # Ekstra vekt = inkl - uten
    ekstra_vekt = vekt_inkl_gen - vekt_uten
    omfar_pst_total = (ekstra_vekt / vekt_uten * 100) if vekt_uten > 0 else 0
    # Vekt-vektet omfar-andel = (inkl-uten)/uten = omfar_pst_total
    gjsnitt_andel = omfar_pst_total
    maks_andel = df[C_ANDEL_GEN].fillna(0).max()

    print(f"\n[INFO] {label}:")
    print(f"  Antall LM-jern:              {ant_total}")
    print(f"  Herav med omfar (L>12m):     {ant_med_omfar}")
    print(f"  Total vekt uten omfar:       {vekt_uten:,.2f} kg")
    print(f"  Total vekt inkl. omfar Gen:  {vekt_inkl_gen:,.2f} kg")
    print(f"  Ekstra vekt pga omfar:       {ekstra_vekt:,.2f} kg  (+{omfar_pst_total:.2f}%)")
    print(f"  Vekt-vektet omfar-andel:     {gjsnitt_andel:.3f}%")
    print(f"  Maks omfar-andel per jern:   {maks_andel:.3f}%")

    return {
        "label":            label,
        "ant_total":        ant_total,
        "ant_med_omfar":    int(ant_med_omfar),
        "vekt_uten":        round(vekt_uten, 2),
        "vekt_inkl_gen":    round(vekt_inkl_gen, 2),
        "ekstra_vekt":      round(ekstra_vekt, 2),
        "omfar_pst_total":  round(omfar_pst_total, 2),
        "gjsnitt_andel":    round(gjsnitt_andel, 3),
        "maks_andel":       round(maks_andel, 3),
    }

s_vek = stats_for(df_vek, "VEK")
s_lya = stats_for(df_lya, "LYA")

# Kombinert total
ant_total_komb   = s_vek["ant_total"]       + s_lya["ant_total"]
ant_omfar_komb   = s_vek["ant_med_omfar"]   + s_lya["ant_med_omfar"]
vekt_uten_komb   = s_vek["vekt_uten"]        + s_lya["vekt_uten"]
vekt_inkl_komb   = s_vek["vekt_inkl_gen"]   + s_lya["vekt_inkl_gen"]
ekstra_komb      = vekt_inkl_komb - vekt_uten_komb
pst_komb         = round(ekstra_komb / vekt_uten_komb * 100, 2) if vekt_uten_komb > 0 else 0

# Vekt-vektet omfar-andel kombinert = pst_komb (identisk med (inkl-uten)/uten)
df_komb_temp = pd.concat([df_vek, df_lya], ignore_index=True)
gjsnitt_komb = pst_komb
maks_komb    = round(df_komb_temp[C_ANDEL_GEN].fillna(0).max(), 3)

print(f"\n[INFO] Kombinert VEK + LYA:")
print(f"  Totalt antall LM-jern:         {ant_total_komb}")
print(f"  Herav med omfar:               {ant_omfar_komb}")
print(f"  Total vekt uten omfar:         {vekt_uten_komb:,.2f} kg")
print(f"  Total vekt inkl. omfar:        {vekt_inkl_komb:,.2f} kg")
print(f"  Ekstra vekt pga omfar:         {ekstra_komb:,.2f} kg  (+{pst_komb:.2f}%)")
print(f"  Vekt-vektet omfar-andel:       {gjsnitt_komb:.3f}%")
print(f"  Maks omfar-andel per jern:     {maks_komb:.3f}%")

# === FORDELING PER DIAMETER ===
def diam_stats(df, label):
    rows = []
    for diam, grp in df.groupby(C_DIAM):
        ant = len(grp)
        ant_omfar = (grp[C_ANT_GEN].fillna(0) > 0).sum()
        vekt_u = grp[C_VEKT].sum()
        vekt_i = grp[C_VEKT_GEN].sum()
        ekstra = vekt_i - vekt_u
        pst    = round(ekstra / vekt_u * 100, 2) if vekt_u > 0 else 0
        # Vekt-vektet andel per diameter = (vekt_inkl-vekt_uten)/vekt_uten = pst
        gjsnitt = pst
        maks_ant = int(grp[C_ANT_GEN].fillna(0).max())
        rows.append({
            "Kilde":                        label,
            "Diameter (mm)":                int(diam),
            "Antall jern":                  ant,
            "Antall med omfar":             int(ant_omfar),
            "Vekt uten omfar (kg)":         round(vekt_u, 2),
            "Vekt inkl. omfar (kg)":        round(vekt_i, 2),
            "Ekstra vekt pga omfar (kg)":   round(ekstra, 2),
            "Omfar-tillegg (%)":            pst,
            "Vekt-vektet omfar-andel (%)":  gjsnitt,
            "Maks antall omfar per jern":   maks_ant,
        })
    return rows

diam_rows = diam_stats(df_vek, "VEK") + diam_stats(df_lya, "LYA")
df_diam = pd.DataFrame(diam_rows)

# === LAG KOMBINERT DF ===
df_vek_komb = df_vek.copy()
df_vek_komb.insert(0, "Kilde", "VEK")
df_lya_komb = df_lya.copy()
df_lya_komb.insert(0, "Kilde", "LYA")
df_kombinert = pd.concat([df_vek_komb, df_lya_komb], ignore_index=True)

# === SKRIV TIL EXCEL ===
print(f"\n[INFO] Skriver til {OUTPUT_FILE} ...")

wb = Workbook()
wb.remove(wb.active)  # fjern default ark

# === STILDEFINISJONER ===
_THIN  = Side(style="thin",   color="BDBDBD")
_BRD   = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_THICK = Side(style="medium", color="9E9E9E")
_BRD_T = Border(left=_THICK, right=_THICK, top=_THICK, bottom=_THICK)

F_TITLE    = Font(bold=True,  name="Calibri", size=16, color="1F497D")
F_H2       = Font(bold=True,  name="Calibri", size=12, color="1F497D")
F_HDR      = Font(bold=True,  name="Calibri", size=10, color="FFFFFF")
F_DATA     = Font(           name="Calibri", size=10)
F_BOLD     = Font(bold=True, name="Calibri", size=10)
F_SMALL    = Font(           name="Calibri", size=9, color="555555")

FILL_BLA   = PatternFill("solid", fgColor="1F497D")  # mork bla
FILL_VEK   = PatternFill("solid", fgColor="2E75B6")  # bla
FILL_LYA   = PatternFill("solid", fgColor="375623")  # gronn
FILL_KOMB  = PatternFill("solid", fgColor="7B3F00")  # brun/oransje
FILL_ZEBRA = PatternFill("solid", fgColor="EBF3FB")  # lys bla
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
FILL_GREY  = PatternFill("solid", fgColor="F5F5F5")
FILL_HDR_DIAM = PatternFill("solid", fgColor="5C5C5C")

AL_CTR  = Alignment(horizontal="center", vertical="center")
AL_L    = Alignment(horizontal="left",   vertical="center")
AL_R    = Alignment(horizontal="right",  vertical="center")
AL_WRAP = Alignment(horizontal="left",   vertical="top", wrap_text=True)

# ============================================================
# ARK 1: Forklaring og Sammendrag
# ============================================================
ws = wb.create_sheet("Forklaring og Sammendrag")
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 20
ws.column_dimensions["D"].width = 20
ws.column_dimensions["E"].width = 22
ws.column_dimensions["F"].width = 22
ws.column_dimensions["G"].width = 18
ws.column_dimensions["H"].width = 22
ws.column_dimensions["I"].width = 22
ws.column_dimensions["J"].width = 12
ws.column_dimensions["K"].width = 14

def _celle(ws, row, col, verdi=None, font=None, fill=None, align=None,
           border=None, fmt=None, bold=False):
    c = ws.cell(row=row, column=col, value=verdi)
    if font:
        c.font = font
    elif bold:
        c.font = Font(bold=True, name="Calibri", size=10)
    else:
        c.font = F_DATA
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    if fmt:
        c.number_format = fmt
    return c

r = 1
# Tittel
_celle(ws, r, 1, "Omfar-analyse LM-armering - VEK og LYA", font=F_TITLE)
r += 1
_celle(ws, r, 1, f"Generert: {datetime.now().strftime('%d.%m.%Y %H:%M')}", font=F_SMALL)
_celle(ws, r, 2, "Kildefiler:", font=F_SMALL)
_celle(ws, r, 3, VEK_FILE.split("\\")[-1], font=F_SMALL)
r += 1
_celle(ws, r, 3, LYA_FILE.split("\\")[-1], font=F_SMALL)
r += 2

# Beskrivelse
_celle(ws, r, 1, "Om dette dokumentet", font=F_H2)
r += 1
beskrivelse = (
    "Dette arket oppsummerer omfar-beregninger for LOSMASSER (LM) armering "
    "fra to entrepriser: VEK (Vegseksjon/tunnelkonstruksjon) og LYA (Lysakerbekken). "
    "Data er hentet fra LM_Omfar-arket i de respektive mengdeuttak-filene. "
    "B35-faktorer (NS 3420): Generelt: 50xD for D<=20mm, 60xD for D>=25mm. "
    "Overkant dekker T>250mm: 60xD for D<=20mm, 80xD for D>=25mm. "
    "Omfar oppstar der jernlengde L>12m (standard stangkleng). "
    "Antall omfar = ceil((L-12) / (12 - omfarlengde))."
)
c_besk = ws.cell(row=r, column=1, value=beskrivelse)
c_besk.font = F_DATA
c_besk.alignment = AL_WRAP
ws.row_dimensions[r].height = 62
r += 2

# ---- SAMMENDRAG-TABELL ----
_celle(ws, r, 1, "Sammendrag - VEK vs LYA vs Total", font=F_H2)
r += 1

hdr_cols = [
    "Kilde",
    "Antall LM-jern totalt",
    "Antall jern med omfar (L>12m)",
    "Total vekt uten omfar (kg)",
    "Total vekt inkl. omfar - Generelt (kg)",
    "Ekstra vekt pga omfar (kg)",
    "Omfar-tillegg totalt (%)",
    "Vekt-vektet omfar-andel (%)",
    "Maks omfar-andel per jern (%)",
]
for ci, h in enumerate(hdr_cols, 1):
    c = ws.cell(row=r, column=ci, value=h)
    c.font = F_HDR
    c.fill = FILL_BLA
    c.alignment = AL_WRAP
    c.border = _BRD
    ws.row_dimensions[r].height = 36

data_rows = [
    ("VEK",
     s_vek["ant_total"], s_vek["ant_med_omfar"],
     s_vek["vekt_uten"], s_vek["vekt_inkl_gen"], s_vek["ekstra_vekt"],
     s_vek["omfar_pst_total"], s_vek["gjsnitt_andel"], s_vek["maks_andel"],
     FILL_ZEBRA),
    ("LYA",
     s_lya["ant_total"], s_lya["ant_med_omfar"],
     s_lya["vekt_uten"], s_lya["vekt_inkl_gen"], s_lya["ekstra_vekt"],
     s_lya["omfar_pst_total"], s_lya["gjsnitt_andel"], s_lya["maks_andel"],
     FILL_WHITE),
    ("VEK + LYA (total)",
     ant_total_komb, ant_omfar_komb,
     vekt_uten_komb, vekt_inkl_komb, ekstra_komb,
     pst_komb, gjsnitt_komb, maks_komb,
     FILL_GREY),
]

r += 1
fmts = [None, "0", "0", "#,##0.00", "#,##0.00", "#,##0.00", '0.00"%"', '0.000"%"', '0.000"%"']
fills_row = [FILL_ZEBRA, FILL_WHITE, FILL_GREY]
for ri, (kilde, a_tot, a_omfar, v_uten, v_inkl, ekstra, pst, gj, maks, fill_r) in enumerate(data_rows):
    vals = [kilde, a_tot, a_omfar, v_uten, v_inkl, ekstra, pst, gj, maks]
    for ci, (v, f) in enumerate(zip(vals, fmts), 1):
        c = ws.cell(row=r + ri, column=ci, value=v)
        c.font = F_BOLD if ci == 1 else F_DATA
        c.fill = fill_r
        c.alignment = AL_R if ci > 1 else AL_L
        c.border = _BRD
        if f:
            c.number_format = f

r += len(data_rows) + 2

# ---- NOKKELTALL VEK ----
def _nokkeltall_seksjon(ws, r, label, s, fill_hdr):
    c = ws.cell(row=r, column=1, value=f"Nokkeltall {label}")
    c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    c.fill = fill_hdr
    c.alignment = AL_L
    c.border = _BRD
    ws.row_dimensions[r].height = 18
    r += 1
    nokkeltall = [
        ("Antall LM-jern",               s["ant_total"],      "0"),
        ("Herav med omfar (L>12m)",      s["ant_med_omfar"],  "0"),
        ("Total vekt uten omfar",        s["vekt_uten"],      '#,##0.00" kg"'),
        ("Total vekt inkl. omfar (Gen.)",s["vekt_inkl_gen"],  '#,##0.00" kg"'),
        ("Ekstra vekt pga omfar",        f"{s['vekt_inkl_gen'] - s['vekt_uten']:,.2f} kg  (+{s['omfar_pst_total']:.2f}%)", None),
        ("Vekt-vektet omfar-andel",       s["gjsnitt_andel"],  '0.000"%"'),
        ("Maks omfar-andel per jern",    s["maks_andel"],     '0.000"%"'),
    ]
    for (navn, verdi, fmt) in nokkeltall:
        c1 = ws.cell(row=r, column=1, value=navn)
        c1.font = F_BOLD
        c1.fill = FILL_GREY
        c1.border = _BRD
        c2 = ws.cell(row=r, column=2, value=verdi)
        c2.font = F_DATA
        c2.fill = FILL_WHITE
        c2.alignment = AL_R
        c2.border = _BRD
        if fmt:
            c2.number_format = fmt
        r += 1
    return r + 1

r = _nokkeltall_seksjon(ws, r, "VEK",  s_vek, FILL_VEK)
r = _nokkeltall_seksjon(ws, r, "LYA",  s_lya, FILL_LYA)

# Kombinert seksjon
c = ws.cell(row=r, column=1, value="Kombinert gjennomsnitt VEK + LYA")
c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
c.fill = FILL_KOMB
c.alignment = AL_L
c.border = _BRD
ws.row_dimensions[r].height = 18
r += 1
komb_nokkeltall = [
    ("Totalt antall LM-jern",        ant_total_komb,  "0"),
    ("Herav med omfar",              ant_omfar_komb,  "0"),
    ("Total vekt uten omfar",        vekt_uten_komb,  '#,##0.00" kg"'),
    ("Total vekt inkl. omfar",       vekt_inkl_komb,  '#,##0.00" kg"'),
    ("Total ekstra vekt pga omfar",  f"{ekstra_komb:,.2f} kg  (+{pst_komb:.2f}%)", None),
    ("Vekt-vektet omfar-andel",       gjsnitt_komb,    '0.000"%"'),
    ("Maks omfar-andel per jern",    maks_komb,       '0.000"%"'),
]
for (navn, verdi, fmt) in komb_nokkeltall:
    c1 = ws.cell(row=r, column=1, value=navn)
    c1.font = F_BOLD
    c1.fill = FILL_GREY
    c1.border = _BRD
    c2 = ws.cell(row=r, column=2, value=verdi)
    c2.font = F_DATA
    c2.fill = FILL_WHITE
    c2.alignment = AL_R
    c2.border = _BRD
    if fmt:
        c2.number_format = fmt
    r += 1
r += 1

# ---- FORDELING PER DIAMETER ----
c = ws.cell(row=r, column=1, value="Fordeling per diameter")
c.font = F_H2
r += 1

diam_hdrs = list(df_diam.columns)
for ci, h in enumerate(diam_hdrs, 1):
    c = ws.cell(row=r, column=ci, value=h)
    c.font = F_HDR
    c.fill = FILL_HDR_DIAM
    c.alignment = AL_WRAP
    c.border = _BRD
    ws.row_dimensions[r].height = 30

diam_fmts = [None, "0", "0", "0", "#,##0.00", "#,##0.00", "#,##0.00", '0.00"%"', '0.000"%"', "0"]
r += 1
for ri, row_data in enumerate(diam_rows):
    is_vek = row_data["Kilde"] == "VEK"
    fill_r = PatternFill("solid", fgColor="DEEAF1") if is_vek else PatternFill("solid", fgColor="E2EFDA")
    for ci, (col, fmt) in enumerate(zip(diam_hdrs, diam_fmts), 1):
        v = row_data[col]
        c = ws.cell(row=r + ri, column=ci, value=v)
        c.font = F_DATA
        c.fill = fill_r
        c.alignment = AL_R if ci > 2 else AL_L
        c.border = _BRD
        if fmt:
            c.number_format = fmt

ws.freeze_panes = "A2"

# ============================================================
# Helper: skriv en LM_Omfar df til ny worksheet
# ============================================================
OMFAR_FMTS = {
    "Vekt (kg)":                   "#,##0.00",
    "Diameter (mm)":               "0",
    "Lengde (m)":                  "#,##0.000",
    "Antall omfar Generelt [=ceil((L-12)/(12-omfar_L))]": "0",
    "Mengde omfar Generelt (m) [=n*omfar_L]":             "#,##0.000",
    "Omfar andel Generelt (%) [=(vekt_inkl-vekt)/vekt*100]": '0.00"%"',
    "Vekt inkl. omfar Generelt (kg)":                     "#,##0.00",
    "Antall omfar Dekker >250mm":                         "0",
    "Mengde omfar Dekker >250mm (m)":                     "#,##0.000",
    "Omfar andel Dekker >250mm (%)":                      '0.00"%"',
    "Vekt inkl. omfar Dekker >250mm (kg)":                "#,##0.00",
    "Vekt inkl. omfar Valgt (kg) [etter PE-regel]":       "#,##0.00",
    "Svinn siste stav (m) [=(n+1)*12-L-n*omfar_L]":       "#,##0.000",
    "Lengde siste stav (m) [=L-n*(12-omfar_L)]":          "#,##0.000",
    "Omfarlengde Generelt (m) [B35: 50D/D<=20, 60D/D>=25]":      "#,##0.000",
    "Omfarlengde Dekker >250mm (m) [B35: 60D/D<=20, 80D/D>=25]": "#,##0.000",
}

def _skriv_omfar_ark(wb, df, arknavn, hdr_fill, tabell_navn):
    ws = wb.create_sheet(arknavn)
    ws.sheet_view.showGridLines = False
    cols = list(df.columns)
    ncols = len(cols)
    nrows = len(df)

    # Header
    for ci, h in enumerate(cols, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = F_HDR
        c.fill = hdr_fill
        c.alignment = AL_WRAP
        c.border = _BRD

    # Data
    for ri, (_, row_s) in enumerate(df.iterrows(), 2):
        is_zebra = (ri % 2 == 0)
        fill_d = FILL_ZEBRA if is_zebra else FILL_WHITE
        for ci, col in enumerate(cols, 1):
            v = row_s[col]
            if pd.isna(v):
                v = None
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = F_DATA
            c.fill = fill_d
            c.alignment = AL_R if ci > 1 else AL_L
            c.border = _BRD
            fmt = OMFAR_FMTS.get(col)
            if fmt:
                c.number_format = fmt

    # Kolonnebredder
    for ci, col in enumerate(cols, 1):
        cl = get_column_letter(ci)
        mx = max(len(str(col)), 8)
        for ri in range(2, min(nrows + 2, 52)):
            v = ws.cell(row=ri, column=ci).value
            if v is not None:
                mx = max(mx, len(str(v)))
        ws.column_dimensions[cl].width = min(mx + 2, 50)

    # Tabell
    if nrows > 0:
        ref = f"A1:{get_column_letter(ncols)}{nrows + 1}"
        tbl = Table(displayName=tabell_navn, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False, showLastColumn=False,
            showRowStripes=False, showColumnStripes=False,
        )
        ws.add_table(tbl)

    ws.freeze_panes = "B2"
    print(f"  [OK] Ark '{arknavn}': {nrows} rader x {ncols} kolonner")
    return ws


# ============================================================
# ARK 2: VEK_LM_Omfar
# ============================================================
print("[INFO] Skriver VEK_LM_Omfar ...")
_skriv_omfar_ark(wb, df_vek, "VEK_LM_Omfar", FILL_VEK, "VEK_LM_Omfar")

# ============================================================
# ARK 3: LYA_LM_Omfar
# ============================================================
print("[INFO] Skriver LYA_LM_Omfar ...")
_skriv_omfar_ark(wb, df_lya, "LYA_LM_Omfar", FILL_LYA, "LYA_LM_Omfar")

# ============================================================
# ARK 4: Kombinert VEK+LYA
# ============================================================
print("[INFO] Skriver Kombinert VEK+LYA ...")
_skriv_omfar_ark(wb, df_kombinert, "Kombinert VEK+LYA", FILL_KOMB, "Kombinert_VEK_LYA")

# ============================================================
# Lagre
# ============================================================
wb.save(OUTPUT_FILE)
print(f"\n[FERDIG] Lagret: {OUTPUT_FILE}")
