from openpyxl import load_workbook

omega_path = r"C:\Users\ifi\Scan Survey\6900 - Fornebubanen\6900-01-010 K7 - Lysaker og Vækerø Stasjon\APM FOB Gjennomgang\Omega_K7_April.xlsx"
wb = load_workbook(omega_path, read_only=True, data_only=True)
ark = wb["BOQNT"] if "BOQNT" in wb.sheetnames else wb.active

# Skriv ut dei 6 første radane for å finna header
rows = []
for rad in ark.iter_rows(values_only=True):
    rows.append(rad)
    if len(rows) >= 6:
        break

for i, rad in enumerate(rows, 1):
    print(f"Rad {i}: {[str(c or '')[:25] for c in rad[:12]]}")

print()

# Finn header-rada (der 'prosesskode' eller 'post' finst)
wb2 = load_workbook(omega_path, read_only=True, data_only=True)
ark2 = wb2["BOQNT"] if "BOQNT" in wb2.sheetnames else wb2.active
headers = None
header_rad_nr = None
for rad_nr, rad in enumerate(ark2.iter_rows(values_only=True), 1):
    celler = [str(c or "").strip().lower() for c in rad]
    if any("prosesskode" in c or "post" == c for c in celler):
        headers = list(rad)
        header_rad_nr = rad_nr
        print(f"Header funne paa rad {rad_nr}")
        print("Kolonnar:", [str(h or "")[:30] for h in headers])
        print()

        # Finn "sist"-kolonnar
        print("=== Kolonnar med 'sist', 'dato', 'brev', 'oppdatert' ===")
        for i, h in enumerate(headers, 1):
            val = str(h or "").lower()
            if any(s in val for s in ["sist", "dato", "brev", "oppdatert"]):
                print(f"  Kol {i}: '{h}'")
        break

# Finn verdiane for den aktuelle posten
wb3 = load_workbook(omega_path, read_only=True, data_only=True)
ark3 = wb3["BOQNT"] if "BOQNT" in wb3.sheetnames else wb3.active
post_kol = None
if headers:
    for i, h in enumerate(headers):
        if "prosesskode" in str(h or "").lower() or str(h or "").lower() == "post":
            post_kol = i
            break

print()
print(f"=== Verdiar for 1302.01.05.241.14.2 (post_kol={post_kol}) ===")
for rad_nr, rad in enumerate(ark3.iter_rows(values_only=True), 1):
    if rad_nr <= (header_rad_nr or 1):
        continue
    if post_kol is not None and str(rad[post_kol] or "").strip() == "1302.01.05.241.14.2":
        for i, (h, v) in enumerate(zip(headers, rad), 1):
            if v is not None:
                print(f"  [{i}] {h}: {v!r}")
        break

wb3.close()
