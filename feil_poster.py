# -*- coding: utf-8 -*-
"""Liste over mengdelisteposter som er feil — éi rad per feil per post.

Stega i køyringa melder inn det dei finn, og til slutt vert alt skrive til
`feil_mengdelisteposter_YYYYMMDD_HHMMSS.xlsx` i utmappa.

Feiltypar:
    Mangler post        Elementet i modellen har ingen FOB_Mengdelistepost.
                        Av som standard — sjå TA_MED_MANGLER_POST.
    Feil materiale      Materialet på elementet skil seg frå det dei fleste
                        elementa på same post har (csv_til_armrib).
    Feil enhet          APM-enheten skil seg frå Omega-enheten, eller modellen
                        har målt i ei anna eining enn posten vert gjord opp i.
    Kun foreldrekode    Posten i modellen er ein foreldrekode — Omega har
                        underpostar, men ingen av dei står på elementet.
    Feil i postnummer   Posten i modellen finst ikkje i Omega, men ei retta
                        utgåve gjer det (manglande '1' først, '030' -> '03',
                        ekstra siffer først, '.' <-> '-'). Mengdeuttaket brukar
                        framleis posten slik han står.
    Rettet automatisk   Posten slik han står i APM-PDF-en er retta av koden
                        (t.d. manglande '1' først, '030' -> '03', '.' -> '-').
                        Vert òg lagra i arket "Rettede poster" i
                        Maalebrev_samlet.xlsx, så dei kjem med når APM-steget
                        er hoppa over.

"Finnes ikke i Omega" er ikkje ein feil og vert ikkje teke med.

Same feil på same post vert slått saman til éi rad, uansett kor mange element
eller APM-ar han gjeld. Rada viser kor mange element det er, og kva
produksjonseiningar og modellar/filer dei ligg i.

Bruk:
    import feil_poster
    feil_poster.nullstill()                       # ved start av køyringa
    feil_poster.registrer("Mangler post", kilde="Modell", element=guid, ...)
    feil_poster.skriv_excel(utmappe)              # til slutt
"""

from __future__ import annotations

import os
import re
from datetime import datetime

MANGLER_POST = "Mangler post"
FEIL_MATERIALE = "Feil materiale"
FEIL_ENHET = "Feil enhet"
RETTET_AUTOMATISK = "Rettet automatisk"
KUN_FORELDREKODE = "Kun foreldrekode"
FEIL_POSTNUMMER = "Feil i postnummer"

FEILTYPAR = (KUN_FORELDREKODE, FEIL_POSTNUMMER, MANGLER_POST, FEIL_MATERIALE,
             FEIL_ENHET, RETTET_AUTOMATISK)

# Element utan mengdelistepost er ofte heilt i orden (dører, utstyr …) og
# fløymde lista. Slå på att her om dei skal vere med.
TA_MED_MANGLER_POST = False

# Ei rad per (feiltype, kilde, post, rettet til, detalj). Element, PE og
# modell vert samla opp i sett på rada.
_NOKKEL = ("feiltype", "kilde", "post", "rettet_til", "detalj")
_SAMLA = ("element", "pe", "modell")

# (overskrift, nøkkel, breidd)
KOLONNER = [
    ("Feiltype", "feiltype", 18),
    ("Kilde", "kilde", 10),
    ("Mengdelistepost", "post", 26),
    ("Rettet til", "rettet_til", 26),
    ("Antall elementer / APM", "antall", 12),
    ("Element / APM", "element", 40),
    ("Produksjonsenheter", "pe", 30),
    ("Modell / fil", "modell", 40),
    ("Detalj", "detalj", 60),
]
MAKS_LISTA = 10      # så mange element/PE/modellar vert lista før "… (+N)"

_FEIL: dict[tuple, dict] = {}


def nullstill() -> None:
    """Tøm lista ved start av ei ny køyring."""
    _FEIL.clear()


def _tekst(v) -> str:
    s = "" if v is None else str(v).strip()
    return "" if s.lower() in ("nan", "none", "<na>") else s


def registrer(feiltype: str, *, kilde: str = "", post="", rettet_til="",
              element="", pe="", modell="", detalj="") -> None:
    """Meld inn eitt tilfelle. Same feil på same post vert slått saman."""
    if feiltype == MANGLER_POST and not TA_MED_MANGLER_POST:
        return
    felt = {"feiltype": feiltype, "kilde": kilde, "post": _tekst(post),
            "rettet_til": _tekst(rettet_til), "detalj": _tekst(detalj)}
    nokkel = tuple(felt[k] for k in _NOKKEL)
    rad = _FEIL.setdefault(nokkel, {**felt, **{k: set() for k in _SAMLA}})
    if _tekst(element):
        rad["element"].add(_tekst(element))
    if _tekst(modell):
        rad["modell"].add(_tekst(modell))
    # PE-feltet kan ha fleire einingar skilde med komma
    rad["pe"].update(p.strip() for p in _tekst(pe).split(",") if p.strip())


def antal() -> int:
    return len(_FEIL)


def _liste(verdiar: set) -> str:
    s = sorted(verdiar)
    if len(s) <= MAKS_LISTA:
        return ", ".join(s)
    return ", ".join(s[:MAKS_LISTA]) + f" … (+{len(s) - MAKS_LISTA})"


# ── Sjekk av mengdelistepostane i modellen ───────────────────────────────────

def post_varianter(post: str) -> list[str]:
    """Retta utgåver av ein post, i prøverekkjefølgje. Posten sjølv er ikkje med.

    Same rettingar som vert gjorde på postane i APM-PDF-ane:
      - '1' manglar først          303-03.87.1    -> 1303-03.87.1
      - ekstra siffer først        11302-03.81.1  -> 1302-03.81.1
      - 0 for mykje i andre ledd   1303-030.87.1  -> 1303-03.87.1
                                   1303-100.87.1  -> 1303-10.87.1
      - første skiljeteikn         1303.03.82.191 <-> 1303-03.82.191
      - '-' seinare i koden        1303-10-43.121 -> 1303-10.43.121
    """
    start = [post]
    if not post.startswith("1"):
        start.append("1" + post)
    forste = re.match(r"^\d+", post)
    if forste and len(forste.group()) == 5:
        start.append(post[1:])

    ledd = []
    for s in start:
        ledd.append(s)
        m = re.match(r"^(\d{4}[.-])(\d{2,}0)([.-].*)?$", s)
        if m:
            ledd.append(m.group(1) + m.group(2)[:-1] + (m.group(3) or ""))

    ut = []
    for v in ledd:
        former = [v]
        if "-" in v[5:]:
            former.append(v[:5] + v[5:].replace("-", "."))
        for w in former:
            ut.append(w)
            if re.match(r"^\d{4}\.", w):
                ut.append(w[:4] + "-" + w[5:])
            elif re.match(r"^\d{4}-", w):
                ut.append(w[:4] + "." + w[5:])
    return [v for v in dict.fromkeys(ut) if v != post]


def sjekk_modellposter(df, omega_koder: set[str]) -> None:
    """Finn postar i modellen som er feil, og meld dei inn per element.

    df: CSV-ane frå csv_til_armrib.les_csv (kolonnenamn med '.').
    omega_koder: alle prosesskodane i Omega.

    Kun foreldrekode: Omega har underpostar til koden, og ingen av dei står i
    same celle. Feil i postnummer: koden finst ikkje i Omega, men ei utgåve frå
    post_varianter gjer det. Kodar som ikkje finst i Omega i det heile er ikkje
    feil, og vert ikkje melde.
    """
    meng_kol = "FOB_Info.FOB_Mengdelistepost"
    if df is None or meng_kol not in df.columns or not omega_koder:
        return

    foreldre: set[str] = set()
    for k in omega_koder:
        delar = re.split(r"([.-])", k)
        for i in range(1, len(delar) - 1, 2):
            foreldre.add("".join(delar[:i]))

    def _finst(k: str) -> bool:
        return k in omega_koder or k in foreldre

    # Same celle går igjen på mange element — rekn ut kvar celle éin gong
    funn_per_celle: dict[str, list[tuple]] = {}
    for celle in df[meng_kol].dropna().astype(str).unique():
        koder = [k.strip() for k in celle.split(";") if k.strip()]
        funn = []
        for k in koder:
            if k in foreldre:
                har_barn = any(o != k and o.startswith((k + ".", k + "-"))
                               for o in koder)
                if not har_barn:
                    funn.append((KUN_FORELDREKODE, k, "",
                                 "Omega har underposter til denne koden — "
                                 "elementet bør ha underposten"))
            elif k not in omega_koder:
                rett = next((v for v in post_varianter(k) if _finst(v)), "")
                if rett:
                    funn.append((FEIL_POSTNUMMER, k, rett,
                                 "Finnes ikke i Omega, men rettet versjon gjør det"))
        if funn:
            funn_per_celle[celle] = funn

    if not funn_per_celle:
        return

    def _f(rad, kol):
        return rad[kol] if kol in rad.index else ""

    from csv_til_armrib import modellnavn
    n = 0
    treff = df[df[meng_kol].astype(str).isin(funn_per_celle.keys())]
    for _, rad in treff.iterrows():
        for feiltype, post, rett, detalj in funn_per_celle[str(rad[meng_kol])]:
            registrer(feiltype, kilde="Modell", post=post, rettet_til=rett,
                      element=_f(rad, "GUID"),
                      pe=_f(rad, "FOB_Info.FOB_Produksjonsenhet"),
                      modell=modellnavn(_f(rad, "Path")), detalj=detalj)
            n += 1
    kun = sum(1 for f in funn_per_celle.values() for x in f if x[0] == KUN_FORELDREKODE)
    feil = sum(1 for f in funn_per_celle.values() for x in f if x[0] == FEIL_POSTNUMMER)
    print(f"[INFO] Modellpostar: {kun} kodar er berre foreldrekode, {feil} har feil "
          f"postnummer — {n:,} element lagt i feillista")


def skriv_excel(utmappe: str) -> str:
    """Skriv lista til utmappa. Returnerer stien, eller "" når lista er tom."""
    if not _FEIL:
        print("[INFO] Ingen feil mengdelisteposter funne — ingen feilliste skriven")
        return ""

    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from lagring import vent_til_skrivbar

    typar = tuple(ft for ft in FEILTYPAR
                  if ft != MANGLER_POST or TA_MED_MANGLER_POST)
    rader = [{**{k: r[k] for k in _NOKKEL},
              "antall": len(r["element"]) or None,
              **{k: _liste(r[k]) for k in _SAMLA}}
             for r in _FEIL.values()]
    rader.sort(key=lambda r: (
        FEILTYPAR.index(r["feiltype"]) if r["feiltype"] in FEILTYPAR else 99,
        r["post"], r["detalj"]))

    wb = openpyxl.Workbook()
    ws_sum = wb.active
    ws_sum.title = "Sammendrag"
    ws = wb.create_sheet("Feil mengdelisteposter")

    # ── Sammendrag ───────────────────────────────────────────────────────────
    ws_sum["A1"] = "Feil mengdelisteposter"
    ws_sum["A1"].font = Font(bold=True, size=14)
    ws_sum["A2"] = f"Køyrd {datetime.now():%d.%m.%Y %H:%M}"
    ws_sum["A4"], ws_sum["B4"] = "Feiltype", "Antall poster"
    ws_sum["A4"].font = ws_sum["B4"].font = Font(bold=True)
    r = 5
    for ft in typar:
        ws_sum.cell(row=r, column=1, value=ft)
        ws_sum.cell(row=r, column=2, value=sum(1 for x in rader if x["feiltype"] == ft))
        r += 1
    ws_sum.cell(row=r, column=1, value="Totalt").font = Font(bold=True)
    ws_sum.cell(row=r, column=2, value=len(rader)).font = Font(bold=True)
    ws_sum.cell(row=r + 2, column=1,
                value="«Finnes ikke i Omega» er ikke regnet som feil.").font = Font(
                    italic=True, color="808080")
    ws_sum.column_dimensions["A"].width = 22
    ws_sum.column_dimensions["B"].width = 14

    # ── Liste ────────────────────────────────────────────────────────────────
    for j, (hdr, _, breidd) in enumerate(KOLONNER, 1):
        ws.cell(row=1, column=j, value=hdr).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(j)].width = breidd
    for i, rad in enumerate(rader, 2):
        for j, (_, nokkel, _) in enumerate(KOLONNER, 1):
            c = ws.cell(row=i, column=j, value=rad[nokkel] or None)
            c.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    ref = f"A1:{get_column_letter(len(KOLONNER))}{len(rader) + 1}"
    tbl = Table(displayName="Feil_mengdelisteposter", ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)

    os.makedirs(utmappe, exist_ok=True)
    utfil = os.path.join(utmappe,
                         f"feil_mengdelisteposter_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    vent_til_skrivbar(utfil)
    wb.save(utfil)
    teljing = ", ".join(f"{ft} {sum(1 for x in rader if x['feiltype'] == ft):,}"
                        for ft in typar)
    print(f"[OK] Feilliste: {len(rader):,} feil ({teljing})")
    print(f"     {utfil}")
    return utfil
